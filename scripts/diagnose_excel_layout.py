from __future__ import annotations
import io
import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

WEEKS = [
    (2025, 14, "https://www.pref.niigata.lg.jp/sec/kanyaku/shuho0714.html"),
    (2025, 15, "https://www.pref.niigata.lg.jp/sec/kanyaku/shuho0715.html"),
    (2025, 26, "https://www.pref.niigata.lg.jp/sec/kanyaku/shuho0726.html"),
    (2025, 36, "https://www.pref.niigata.lg.jp/sec/kanyaku/shuho0736.html"),
]

S = requests.Session()
S.headers.update({
    "User-Agent": "Niigata-Infection-Radar-Diagnostic/1.0",
    "Accept-Language": "ja,en-US;q=0.7,en;q=0.3",
})

def norm(v):
    if v is None:
        return ""
    return str(v).replace("\n"," ").replace("\r"," ").strip()

def get(url, binary=False):
    r=S.get(url, timeout=40)
    r.raise_for_status()
    if binary:
        return r.content
    raw = r.content
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError:
        return raw.decode(r.apparent_encoding or "cp932", errors="replace")

def excel_url(page_url):
    soup=BeautifulSoup(get(page_url), "html.parser")
    for a in soup.find_all("a", href=True):
        text=norm(a.get_text(" ",strip=True))
        if "5類感染症定点把握対象疾患報告数" in text:
            from urllib.parse import urljoin
            return urljoin(page_url, a["href"])
    raise RuntimeError("Excel link not found")

def row_dump(ws, r, max_col=36):
    vals=[]
    for c in range(1, min(ws.max_column,max_col)+1):
        v=ws.cell(r,c).value
        if v not in (None,""):
            vals.append(f"{get_column_letter(c)}{r}={v!r}")
    return " | ".join(vals)

def merged_for_rows(ws, start, end):
    out=[]
    for rng in ws.merged_cells.ranges:
        if rng.max_row >= start and rng.min_row <= end:
            out.append(str(rng))
    return out

for year,week,page_url in WEEKS:
    print("\n" + "="*110)
    print(f"{year} W{week:02d}")
    print("PAGE:", page_url)
    xurl=excel_url(page_url)
    print("XLSX:", xurl)

    data=get(xurl,binary=True)
    print("BYTES:", len(data))
    wb=load_workbook(io.BytesIO(data),data_only=False,read_only=False)

    for ws in wb.worksheets:
        print("\n" + "-"*90)
        print(f"SHEET: {ws.title!r} rows={ws.max_row} cols={ws.max_column}")

        # インフルエンザ関連セルを全部探す
        hits=[]
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value,str) and "インフルエンザ" in cell.value:
                    hits.append((cell.row,cell.column,cell.coordinate,cell.value))

        if not hits:
            print("NO influenza text")
            continue

        print("INFLUENZA CELLS:")
        for h in hits[:30]:
            print(" ", h)

        # 各ヒット周辺を広めに出す
        printed=set()
        for rr,cc,coord,val in hits[:12]:
            start=max(1,rr-3)
            end=min(ws.max_row,rr+8)
            key=(start,end)
            if key in printed:
                continue
            printed.add(key)

            print(f"\nAROUND {coord}={val!r}  rows {start}-{end}")
            merges=merged_for_rows(ws,start,end)
            print("MERGES:", ", ".join(merges[:80]) if merges else "(none)")
            for r in range(start,end+1):
                s=row_dump(ws,r,36)
                if s:
                    print(s)

        # 先頭20行も必ず出す。第15週のレイアウト変更確認用
        print("\nTOP 20 ROWS:")
        for r in range(1,min(ws.max_row,20)+1):
            s=row_dump(ws,r,36)
            if s:
                print(s)
