from __future__ import annotations
import io, json, re
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any
import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook

INDEX_URL = "https://www.pref.niigata.lg.jp/sec/kanyaku/1232482573101.html"
UA = "Niigata-Infection-Radar/2.2 (public-data visualization)"
REGIONS = ["新潟市","新発田","新津","三条","長岡","魚沼","南魚沼","十日町","柏崎","糸魚川","村上","佐渡","上越"]
AGE_GROUPS = ["0歳","1～4歳","5～9歳","10～14歳","15～19歳","20～59歳","60歳以上"]

SESSION = requests.Session()
SESSION.headers.update({
    "User-Agent": UA,
    "Accept-Language":"ja,en-US;q=0.7,en;q=0.3"
})

def norm(v: Any) -> str:
    if v is None:
        return ""
    s = str(v).replace("\u3000"," ").replace("\n"," ").replace("\r"," ")
    s = re.sub(r"\s+", "", s)
    s = s.replace("〜","～").replace("－","-").replace("―","-").replace("−","-")
    s = re.sub(r"[※＊*]+$", "", s)
    return s

def get(url: str, binary=False):
    r = SESSION.get(url, timeout=35)
    r.raise_for_status()
    if binary:
        return r.content
    raw = r.content
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError:
        return raw.decode(r.apparent_encoding or "cp932", errors="replace")

def abs_url(base, href):
    from urllib.parse import urljoin
    return urljoin(base, href)

@dataclass
class WeekPage:
    year:int
    week:int
    url:str

def discover_week_pages(start_year=2025):
    soup=BeautifulSoup(get(INDEX_URL),"html.parser")
    out=[]; seen=set()
    for a in soup.find_all("a",href=True):
        txt=a.get_text(" ",strip=True)
        m=re.search(r"令和\s*(\d+)\s*年\s*第\s*(\d+)\s*週",txt)
        if not m:
            continue
        y=2018+int(m.group(1)); w=int(m.group(2))
        if y < start_year:
            continue
        k=(y,w)
        if k not in seen:
            seen.add(k)
            out.append(WeekPage(y,w,abs_url(INDEX_URL,a["href"])))
    out.sort(key=lambda x:(x.year,x.week))
    if not out:
        raise RuntimeError("週報一覧から週ページを検出できませんでした。")
    return out

def parse_dates_from_page(soup, year, week):
    text=soup.get_text(" ",strip=True)
    m=re.search(rf"第\s*{week}\s*週.*?(\d+)月(\d+)日から.*?(\d+)月(\d+)日",text)
    if not m:
        return None,None
    sm,sd,em,ed=map(int,m.groups())
    sy=year; ey=year+(1 if em<sm else 0)
    if week==1 and sm==12:
        sy=year-1; ey=year
    try:
        return date(sy,sm,sd),date(ey,em,ed)
    except ValueError:
        return None,None

def reiwa_label(d1,d2):
    ry=d1.year-2018
    if d1.year==d2.year:
        return f"R{ry}/{d1.month}/{d1.day}–{d2.month}/{d2.day}"
    return f"R{ry}/{d1.month}/{d1.day}–R{d2.year-2018}/{d2.month}/{d2.day}"

def extract_topic(soup):
    text=soup.get_text("\n",strip=True)
    m=re.search(
        r"◆\s*インフルエンザ([\s\S]*?)(?=\n\s*◆|\n\s*警報を発令している疾患|$)",
        text
    )
    if not m:
        return ""
    return re.sub(r"\n{3,}","\n\n","インフルエンザ"+m.group(1)).strip()[:3000]

def extract_topic_prefecture_value(soup):
    topic=extract_topic(soup)
    if not topic:
        return None
    m=re.search(r"全県で\s*([0-9]+(?:\.[0-9]+)?)",topic)
    return float(m.group(1)) if m else None

def find_excel_url(soup,page_url):
    for a in soup.find_all("a",href=True):
        t=norm(a.get_text(" ",strip=True))
        href=a["href"]
        if "5類感染症定点把握対象疾患報告数" in t:
            return abs_url(page_url,href)
        if href.lower().endswith((".xlsx",".xlsm")) and "定点" in t:
            return abs_url(page_url,href)
    return None

def as_number(v):
    if isinstance(v,(int,float)) and not isinstance(v,bool):
        return float(v)
    if isinstance(v,str):
        s=v.strip().replace(",","")
        if re.fullmatch(r"-?\d+(?:\.\d+)?",s):
            return float(s)
    return None

def canonical_header(s):
    s=norm(s)
    if s in ("県計","全県","県全体"):
        return "県計"
    for r in REGIONS:
        if s==r:
            return r
    return None

def find_main_table(ws):
    title_row=None
    for r in range(1, min(ws.max_row, 80)+1):
        for c in range(1, min(ws.max_column, 20)+1):
            if "地域振興局等管内別報告数" in norm(ws.cell(r,c).value):
                title_row=r
                break
        if title_row:
            break
    if title_row is None:
        return None

    header_row=None; cols=None
    for r in range(title_row, min(ws.max_row, title_row+5)+1):
        found={}
        for c in range(1, min(ws.max_column, 30)+1):
            key=canonical_header(ws.cell(r,c).value)
            if key and key not in found:
                found[key]=c
        if len(found) >= 10 and "県計" in found:
            header_row=r; cols=found
            break
    if header_row is None:
        return None

    flu_row=None
    for r in range(header_row+1, min(ws.max_row, header_row+12)+1):
        if norm(ws.cell(r,1).value) == "インフルエンザ":
            flu_row=r
            break
    if flu_row is None:
        return None

    value_row=None
    for r in range(flu_row, min(ws.max_row, flu_row+4)+1):
        left="".join(norm(ws.cell(r,c).value) for c in range(1, min(cols.values())))
        if "定点当" in left:
            value_row=r
            break
    if value_row is None:
        return None

    return {
        "title_row":title_row,
        "header_row":header_row,
        "flu_row":flu_row,
        "value_row":value_row,
        "cols":cols,
    }

# ---------- 年代別 ----------

def parse_age_range_header(value):
    s=norm(value)
    if not s:
        return None
    s=s.replace("年齢","").replace("年令","")
    if s in ("0歳","0"):
        return (0,0)
    m=re.fullmatch(r"(\d+)～(\d+)歳?",s)
    if m:
        return (int(m.group(1)),int(m.group(2)))
    m=re.fullmatch(r"(\d+)-(\d+)歳?",s)
    if m:
        return (int(m.group(1)),int(m.group(2)))
    m=re.fullmatch(r"(\d+)歳?以上",s)
    if m:
        return (int(m.group(1)),None)
    return None

def age_label(age_range):
    lo,hi=age_range
    if hi is None:
        return f"{lo}歳以上"
    if lo==hi:
        return f"{lo}歳"
    return f"{lo}～{hi}歳"

def find_age_table(ws):
    max_r=min(ws.max_row,180)
    max_c=min(ws.max_column,80)
    candidates=[]

    for r in range(1,max_r+1):
        age_cols={}
        for c in range(1,max_c+1):
            ar=parse_age_range_header(ws.cell(r,c).value)
            if ar and ar not in age_cols:
                age_cols[ar]=c
        if len(age_cols) < 5:
            continue

        title_row=None
        title_score=0
        for rr in range(max(1,r-8),r+1):
            rowtext="".join(norm(ws.cell(rr,cc).value) for cc in range(1,min(max_c,20)+1))
            if "年齢" in rowtext:
                title_row=rr
                title_score=20
                break

        flu_row=None
        for rr in range(r+1,min(max_r,r+30)+1):
            left="".join(norm(ws.cell(rr,cc).value) for cc in range(1,min(min(age_cols.values()),10)+1))
            if "インフルエンザ" in left:
                flu_row=rr
                break

        if flu_row is None:
            continue

        candidates.append({
            "score":len(age_cols)*3 + title_score - max(0,flu_row-r-5),
            "title_row":title_row,
            "header_row":r,
            "flu_row":flu_row,
            "age_cols":age_cols,
        })

    if not candidates:
        return None
    candidates.sort(key=lambda x:x["score"],reverse=True)
    return candidates[0]

def row_age_values(ws, row, age_cols):
    vals={}
    numeric=0
    for ar,c in age_cols.items():
        n=as_number(ws.cell(row,c).value)
        if n is not None:
            numeric += 1
            vals[age_label(ar)] = float(n)
        else:
            vals[age_label(ar)] = 0.0
    return vals, numeric

def parse_age_groups(wb, expected_count=None, expected_rate=None):
    """
    年代別表から
      age_counts         = 実数（人）
      age_per_sentinel   = 定点当たり報告数（人/定点）
    の両方を取得する。

    行ラベルだけに依存せず、
      ・実数行は県計実数との合計一致
      ・定点当たり行は県全体定点値との合計近似
    でも判定する。
    """
    diagnostics=[]

    for ws in wb.worksheets:
        info=find_age_table(ws)
        if not info:
            diagnostics.append(f"{ws.title}: 年齢表なし")
            continue

        candidate_rows=[]
        for rr in range(info["flu_row"], min(ws.max_row, info["flu_row"]+5)+1):
            vals,numeric=row_age_values(ws,rr,info["age_cols"])
            if numeric < max(3, len(info["age_cols"])//2):
                continue
            total=sum(vals.values())

            left="".join(
                norm(ws.cell(rr,c).value)
                for c in range(1, min(min(info["age_cols"].values()),10)+1)
            )
            candidate_rows.append({
                "row":rr,
                "vals":vals,
                "total":total,
                "left":left,
            })

        if not candidate_rows:
            diagnostics.append(f"{ws.title}: 年齢表候補に数値行なし")
            continue

        count_row=None
        rate_row=None

        # まずラベルで判定
        for x in candidate_rows:
            if "定点当" in x["left"]:
                rate_row=x
            elif "インフルエンザ" in x["left"] and count_row is None:
                count_row=x

        # 実数は県計実数との合計一致を最優先
        if expected_count is not None:
            best=min(candidate_rows, key=lambda x:abs(x["total"]-float(expected_count)))
            if abs(best["total"]-float(expected_count)) <= 0.01:
                count_row=best

        # 定点当たりは県全体定点値に最も近い合計を選択
        # 各年代を小数2桁に丸めているため、合計には数百分の誤差が出る。
        if expected_rate is not None:
            rate_candidates=[x for x in candidate_rows if count_row is None or x["row"] != count_row["row"]]
            if rate_candidates:
                best=min(rate_candidates, key=lambda x:abs(x["total"]-float(expected_rate)))
                if abs(best["total"]-float(expected_rate)) <= 0.15:
                    rate_row=best

        # 典型レイアウトでは実数行の次行が定点当たり
        if count_row is not None and rate_row is None:
            nxt=next((x for x in candidate_rows if x["row"]==count_row["row"]+1),None)
            if nxt:
                rate_row=nxt

        if count_row is None or rate_row is None:
            diagnostics.append(
                f"{ws.title}: 年代別の実数/定点当たり行を確定できず "
                f"(count={None if count_row is None else count_row['row']}, "
                f"rate={None if rate_row is None else rate_row['row']})"
            )
            continue

        counts={k:round(v,4) for k,v in count_row["vals"].items()}
        rates={k:round(v,4) for k,v in rate_row["vals"].items()}

        count_sum=round(sum(counts.values()),4)
        rate_sum=round(sum(rates.values()),4)

        verification={
            "expected_count":None if expected_count is None else round(float(expected_count),4),
            "age_count_sum":count_sum,
            "count_difference":None if expected_count is None else round(count_sum-float(expected_count),4),
            "expected_rate":None if expected_rate is None else round(float(expected_rate),4),
            "age_rate_sum":rate_sum,
            "rate_difference":None if expected_rate is None else round(rate_sum-float(expected_rate),4),
        }

        if expected_count is not None and abs(count_sum-float(expected_count)) > 0.01:
            diagnostics.append(f"{ws.title}: 年代別実数合計不一致 {count_sum} != {expected_count}")
            continue

        return {
            "age_counts":counts,
            "age_per_sentinel":rates,
            "_age_parser":{
                "sheet":ws.title,
                "title_row":info["title_row"],
                "header_row":info["header_row"],
                "flu_row":info["flu_row"],
                "count_row":count_row["row"],
                "rate_row":rate_row["row"],
                "raw_age_columns":len(info["age_cols"]),
                "verification":verification,
                "strategy":"age_counts_and_per_sentinel"
            }
        }

    return {
        "age_counts":{},
        "age_per_sentinel":{},
        "_age_parser":{
            "strategy":"not_found",
            "diagnostics":diagnostics[:20]
        }
    }

def parse_influenza_excel(data):
    wb=load_workbook(io.BytesIO(data),data_only=True,read_only=False)
    diagnostics=[]
    main_result=None

    for ws in wb.worksheets:
        info=find_main_table(ws)
        if not info:
            diagnostics.append(f"{ws.title}: 主表を特定できず")
            continue

        vals={}
        for key,c in info["cols"].items():
            raw=ws.cell(info["value_row"],c).value
            n=as_number(raw)
            if key=="県計":
                if n is not None:
                    vals[key]=n
            else:
                vals[key]=0.0 if n is None else n

        prefecture_count=None
        if "県計" in info["cols"]:
            prefecture_count=as_number(ws.cell(info["flu_row"],info["cols"]["県計"]).value)

        if "県計" not in vals:
            actual_values=[]
            for key,c in info["cols"].items():
                n_actual=as_number(ws.cell(info["flu_row"],c).value)
                actual_values.append(0.0 if n_actual is None else n_actual)
            if all(v==0.0 for v in actual_values):
                vals["県計"]=0.0
                if prefecture_count is None:
                    prefecture_count=0.0
            else:
                diagnostics.append(f"{ws.title}: 県計なし")
                continue

        region_count=sum(1 for r in REGIONS if r in info["cols"])
        if region_count < 10:
            diagnostics.append(f"{ws.title}: 地域列不足={region_count}")
            continue

        main_result={
            "prefecture":round(float(vals["県計"]),4),
            "prefecture_count":None if prefecture_count is None else round(float(prefecture_count),4),
            "regions":{k:round(float(vals.get(k,0.0)),4) for k in REGIONS if k in info["cols"]},
            "_parser":{
                "sheet":ws.title,
                "title_row":info["title_row"],
                "header_row":info["header_row"],
                "flu_row":info["flu_row"],
                "value_row":info["value_row"],
                "region_count":region_count,
                "strategy":"main_regional_table_only"
            }
        }
        break

    if main_result is None:
        raise ValueError("主表からインフルエンザ定点値を取得できません: "+"; ".join(diagnostics))

    age_result=parse_age_groups(
        wb,
        expected_count=main_result.get("prefecture_count"),
        expected_rate=main_result.get("prefecture")
    )

    main_result.update({
        "age_counts":age_result["age_counts"],
        "age_per_sentinel":age_result["age_per_sentinel"],
        "_age_parser":age_result["_age_parser"],
    })
    return main_result

def scrape_week(wp):
    soup=BeautifulSoup(get(wp.url),"html.parser")
    d1,d2=parse_dates_from_page(soup,wp.year,wp.week)

    xurl=find_excel_url(soup,wp.url)
    if not xurl:
        raise ValueError("Excelリンクが見つかりません")

    parsed=parse_influenza_excel(get(xurl,binary=True))

    topic_value=extract_topic_prefecture_value(soup)
    if topic_value is not None and abs(parsed["prefecture"]-topic_value)>0.011:
        raise ValueError(
            f"本文照合不一致: Excel={parsed['prefecture']} / 今週のトピック={topic_value}"
        )

    return {
        "year":wp.year,
        "week":wp.week,
        "start":d1.isoformat() if d1 else None,
        "end":d2.isoformat() if d2 else None,
        "label":reiwa_label(d1,d2) if d1 and d2 else f"{wp.year}-W{wp.week:02d}",
        "prefecture":parsed["prefecture"],
        "regions":parsed["regions"],
        "age_counts":parsed.get("age_counts",{}),
        "age_per_sentinel":parsed.get("age_per_sentinel",{}),
        "topic":extract_topic(soup),
        "source_page":wp.url,
        "source_excel":xurl,
        "verification":{
            "method":"niigata_main_table" + ("+topic" if topic_value is not None else ""),
            "niigata_topic_value":topic_value,
            "prefecture_count":parsed.get("prefecture_count")
        },
        "_parser":parsed["_parser"],
        "_age_parser":parsed.get("_age_parser",{})
    }

def read_json(path):
    p=Path(path)
    if not p.exists():
        return {"meta":{},"weeks":[]}
    return json.loads(p.read_text(encoding="utf-8"))

def write_json(path,data):
    p=Path(path)
    p.parent.mkdir(parents=True,exist_ok=True)
    p.write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding="utf-8")

def dedupe(weeks):
    d={(w["year"],w["week"]):w for w in weeks}
    return sorted(d.values(),key=lambda x:(x["year"],x["week"]))
